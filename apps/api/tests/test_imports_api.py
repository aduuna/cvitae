"""POST /imports and the uploads it keeps, with the fake providers."""

import io
import zipfile
from pathlib import Path

import docx
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from api.db.models import AiRun
from api.ports.llm.base import ProviderName, provider_failure
from api.ports.llm.fake import FakeLLMProvider
from api.ports.storage import LocalStorage
from cvcore.models import CvContent
from cvcore.render import iter_bullets
from cvcore.wire import WireImportedCv, to_wire

FIXTURES = Path(__file__).resolve().parents[3] / "templates" / "fixtures"

SAMPLE_PDF = (FIXTURES / "sample-cv.pdf").read_bytes()


def docx_bytes(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def runs(session: AsyncSession) -> list[AiRun]:
    return list((await session.scalars(select(AiRun))).all())


async def test_imports_a_pdf_and_keeps_the_file(
    client: AsyncClient,
    mimo: FakeLLMProvider,
    sample_profile: CvContent,
    storage: LocalStorage,
    db_session: AsyncSession,
) -> None:
    response = await client.post(
        "/api/v1/imports", files={"file": ("Ama CV.pdf", SAMPLE_PDF, "application/pdf")}
    )

    assert response.status_code == 200
    body = response.json()
    # MiMo Flash is the default model.
    assert (body["ai_model"], body["provider"], body["model"]) == (
        "mimo-flash",
        "mimo",
        "mimo-v2.6-flash",
    )
    assert body["profile_name"] == "Data engineering"
    draft = CvContent.model_validate(body["draft"])
    # Same facts as the source of sample-cv.pdf, with fresh ids.
    assert draft.basics == sample_profile.basics
    assert [b.text for b in iter_bullets(draft)] == [b.text for b in iter_bullets(sample_profile)]
    assert {b.id for b in iter_bullets(draft)}.isdisjoint(
        b.id for b in iter_bullets(sample_profile)
    )
    [sent] = mimo.requests
    assert sent.document == SAMPLE_PDF
    assert sent.user_text == "Extract the career data from the attached CV."
    uploads = (await client.get("/api/v1/uploads")).json()
    assert [(u["id"], u["filename"], u["mime_type"]) for u in uploads] == [
        (body["upload_id"], "Ama CV.pdf", "application/pdf")
    ]
    assert [(r.task, r.status) for r in await runs(db_session)] == [("import", "succeeded")]


async def test_imports_docx_as_text(client: AsyncClient, mimo: FakeLLMProvider) -> None:
    data = docx_bytes("Ama Owusu-Mensah", "Data engineer", table=[["Python", "Spark"]])

    response = await client.post("/api/v1/imports", files={"file": ("cv.docx", data, "x/y")})

    assert response.status_code == 200
    [sent] = mimo.requests
    assert sent.document is None
    assert "<cv>\nAma Owusu-Mensah\nData engineer\nPython | Spark\n</cv>" in sent.user_text


async def test_the_headline_names_the_profile_when_the_model_does_not(
    client: AsyncClient, mimo: FakeLLMProvider, sample_profile: CvContent
) -> None:
    sample_profile.basics.headline = "  Senior Data Engineer  "
    mimo.responders["import"] = lambda _: WireImportedCv(
        content=to_wire(sample_profile), profile_name=" "
    )

    response = await client.post("/api/v1/imports", data={"text": "Ama, data engineer"})

    assert response.json()["profile_name"] == "Senior Data Engineer"


async def test_imports_pasted_text(client: AsyncClient, mimo: FakeLLMProvider) -> None:
    response = await client.post("/api/v1/imports", data={"text": "  Ama, data engineer  "})

    assert response.status_code == 200
    assert response.json()["upload_id"] is None
    assert "<cv>\nAma, data engineer\n</cv>" in mimo.requests[0].user_text


async def test_imports_an_earlier_upload_again_with_another_provider(
    client: AsyncClient, gemini: FakeLLMProvider
) -> None:
    first = await client.post(
        "/api/v1/imports", files={"file": ("cv.pdf", SAMPLE_PDF, "application/pdf")}
    )
    upload_id = first.json()["upload_id"]

    again = await client.post(
        "/api/v1/imports", data={"upload_id": upload_id, "ai_model": "gemini-flash-lite"}
    )

    assert again.status_code == 200
    assert (again.json()["upload_id"], again.json()["ai_model"]) == (
        upload_id,
        "gemini-flash-lite",
    )
    assert gemini.requests[0].document == SAMPLE_PDF
    assert gemini.requests[0].model == "gemini-3.5-flash-lite"
    assert len((await client.get("/api/v1/uploads")).json()) == 1


async def test_the_users_model_is_the_default(client: AsyncClient, claude: FakeLLMProvider) -> None:
    await client.patch("/api/v1/me", json={"preferred_ai_model": "claude-haiku"})

    response = await client.post("/api/v1/imports", data={"text": "My CV"})

    assert response.json()["ai_model"] == "claude-haiku"
    assert [r.model for r in claude.requests] == ["claude-haiku-5-5"]


async def test_needs_exactly_one_source(client: AsyncClient) -> None:
    none = await client.post("/api/v1/imports", data={})
    both = await client.post(
        "/api/v1/imports",
        data={"text": "My CV"},
        files={"file": ("cv.pdf", SAMPLE_PDF, "application/pdf")},
    )

    for response in (none, both):
        assert response.status_code == 422
        assert response.json()["error"]["message"].startswith("Upload a file, paste your CV")


async def test_refuses_other_files_and_bad_text(client: AsyncClient) -> None:
    html = await client.post(
        "/api/v1/imports", files={"file": ("cv.pdf", b"<html>", "application/pdf")}
    )
    blank = await client.post("/api/v1/imports", data={"text": "   "})
    long = await client.post("/api/v1/imports", data={"text": "x" * 60_001})
    broken_docx = await client.post(
        "/api/v1/imports", files={"file": ("cv.docx", b"PK\x03\x04broken", "x/y")}
    )

    assert html.json()["error"]["message"] == (
        "This file is not a PDF or DOCX. Upload a PDF or DOCX file."
    )
    assert blank.json()["error"]["message"].startswith("There is no text to import")
    assert long.json()["error"]["message"].startswith("This CV is longer than 60,000 characters")
    assert broken_docx.status_code == 422


async def test_unreadable_docx_is_explained(client: AsyncClient) -> None:
    # A zip with word/document.xml that is not a real DOCX.
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "not xml")
    response = await client.post(
        "/api/v1/imports", files={"file": ("cv.docx", buffer.getvalue(), "x/y")}
    )

    assert response.status_code == 422
    assert response.json()["error"]["message"].startswith("This DOCX could not be read")


async def test_provider_failures_are_502_and_still_recorded(
    client: AsyncClient, mimo: FakeLLMProvider, db_session: AsyncSession
) -> None:
    mimo.fail_with = provider_failure(ProviderName.mimo, "rate_limited")

    response = await client.post(
        "/api/v1/imports", files={"file": ("cv.pdf", SAMPLE_PDF, "application/pdf")}
    )

    assert response.status_code == 502
    assert response.json()["error"]["details"] == {
        "provider": "mimo",
        "error_code": "rate_limited",
    }
    assert [(r.task, r.status, r.error_code) for r in await runs(db_session)] == [
        ("import", "failed", "rate_limited")
    ]
    # The file is kept, so the user can try again with the other provider.
    assert len((await client.get("/api/v1/uploads")).json()) == 1


async def test_uploads_are_private_and_deletable(
    client: AsyncClient, other_client: AsyncClient, storage: LocalStorage
) -> None:
    created = await client.post(
        "/api/v1/imports", files={"file": ("cv.pdf", SAMPLE_PDF, "application/pdf")}
    )
    upload_id = created.json()["upload_id"]

    assert (await other_client.get("/api/v1/uploads")).json() == []
    assert (await other_client.delete(f"/api/v1/uploads/{upload_id}")).status_code == 404
    assert (
        await other_client.post("/api/v1/imports", data={"upload_id": upload_id})
    ).status_code == 404

    assert (await client.delete(f"/api/v1/uploads/{upload_id}")).status_code == 204
    assert (await client.get("/api/v1/uploads")).json() == []
