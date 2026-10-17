"""Import: extract a profile draft from a PDF, a DOCX or pasted text. Nothing is saved as a
profile here; the user reviews the draft and saves it with POST /profiles."""

import io
import uuid
import zipfile
from dataclasses import dataclass

import docx

from api.ai import tasks
from api.billing.plans import Plan
from api.errors import ValidationFailed
from api.ports.storage import Storage
from api.repositories.ai_runs import AiRunRepository
from api.repositories.profiles import ProfileRepository
from api.repositories.uploads import UploadRepository
from api.services.ai_models import ModelChoice
from api.services.quotas import require_ai_runs, require_room
from api.services.uploads import PDF, get_upload, store_upload
from cvcore.ids import with_fresh_ids
from cvcore.models import CvContent

MAX_TEXT_CHARS = 60_000


@dataclass(frozen=True)
class ImportDraft:
    upload_id: uuid.UUID | None
    draft: CvContent
    profile_name: str
    """The model's suggestion, else the CV's headline; empty when neither says anything."""
    choice: ModelChoice
    model: str
    """The model id that answered."""


def docx_text(data: bytes) -> str:
    """The text of a DOCX: paragraphs, then table cells, in document order per block."""
    try:
        document = docx.Document(io.BytesIO(data))
    except (KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise ValidationFailed(
            "This DOCX could not be read. Save it again from your editor, or upload a PDF."
        ) from exc
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


async def import_profile(
    choice: ModelChoice,
    runs: AiRunRepository,
    uploads: UploadRepository,
    storage: Storage,
    profiles: ProfileRepository,
    *,
    plan: Plan,
    file: tuple[str, bytes] | None = None,
    text: str | None = None,
    upload_id: uuid.UUID | None = None,
) -> ImportDraft:
    """Exactly one of `file` (name and bytes), pasted `text`, or an earlier `upload_id`."""
    if sum(source is not None for source in (file, text, upload_id)) != 1:
        raise ValidationFailed("Upload a file, paste your CV as text, or choose an earlier upload.")
    # The draft becomes a profile, so check there is room for one before spending a run.
    await require_room("profiles", plan, profiles)
    await require_ai_runs(runs, plan)

    upload = None
    data = b""
    pdf: bytes | None = None
    body: str | None = None
    if file is not None:
        upload = await store_upload(uploads, storage, *file)
        data = file[1]
    elif upload_id is not None:
        upload = await get_upload(uploads, upload_id)
        data = await storage.get(upload.storage_key)
    if upload is not None:
        if upload.mime_type == PDF:
            pdf = data
        else:
            body = docx_text(data)
    else:
        body = (text or "").strip()

    if body is not None:
        if not body:
            raise ValidationFailed("There is no text to import. Paste your CV and try again.")
        if len(body) > MAX_TEXT_CHARS:
            raise ValidationFailed(
                f"This CV is longer than {MAX_TEXT_CHARS:,} characters. Shorten it and try again."
            )

    result = await tasks.extract_profile(choice, runs, text=body, pdf=pdf)
    content = result.output.content
    return ImportDraft(
        upload_id=upload.id if upload else None,
        draft=with_fresh_ids(content),
        profile_name=result.output.profile_name.strip() or content.basics.headline.strip(),
        choice=choice,
        model=result.model,
    )
